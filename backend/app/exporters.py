"""Export a Log Analyzer report to an Excel workbook (openpyxl)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_FILL = PatternFill("solid", fgColor="1F4E79")
_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = _FILL, _FONT
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + c)].width = max(14, len(str(h)) + 2)
    return ws


def to_excel(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Log Analyzer — Summary"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    for i, (k, v) in enumerate(report["summary"].items(), start=3):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    for i, note in enumerate(report["insights"], start=12):
        ws.cell(row=i, column=1, value=f"• {note}")

    _sheet(wb, "Status", ["Class", "Count"], [[r["status_class"], r["count"]] for r in report["by_status_class"]])
    _sheet(wb, "Top Paths", ["Path", "Hits", "Bytes"], [[r["path"], r["hits"], r["bytes"]] for r in report["top_paths"]])
    _sheet(wb, "Top IPs", ["IP", "Requests"], [[r["ip"], r["requests"]] for r in report["top_ips"]])
    _sheet(wb, "Errors", ["Path", "Status", "Count"], [[r["path"], r["status"], r["count"]] for r in report["top_errors"]])
    if report["anomalous_ips"]:
        _sheet(wb, "Anomalous IPs", ["IP", "Requests", "Z-score"],
               [[r["ip"], r["requests"], r["z_score"]] for r in report["anomalous_ips"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
